"""Versandfertige Dokumente: PDF und Word aus Markdown/HTML, und deren Prüfung (#893).

Kommandozeile: ``dokument`` (``agent/scripts/dokument``, im Abbild unter
/usr/local/bin). Alle drei Laufzeiten — Claude Code, Codex, Custom-LLM — rufen
dasselbe Werkzeug über ihre Shell auf; es gibt keine zweite Umsetzung.

    dokument pdf  <eingabe.md|.html> [-o ziel.pdf] [--titel T] [--fusszeile T]
    dokument docx <eingabe.md> [-o ziel.docx]
    dokument pruefen [--json] <datei.pdf>

Warum es das gibt: Ohne Werkzeug druckten die Agenten mit ``chromium
--print-to-pdf``. Das setzt die Standard-Kopfzeile des Browsers (US-Datum,
Titel) und die Fußzeile (``file:///workspace/…``, Seitenzahl „1/3“) auf jede
Seite — so ging kein Angebot an einen Kunden.

``pruefe_pdf`` ist zugleich die Prüfung, mit der ``present_file`` ein solches
PDF abweist: der Executor ruft sie direkt, der MCP-Server über
``dokument pruefen --json``.
"""

from __future__ import annotations

import argparse
import html
import json
import os
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

# ─── Prüfung ──────────────────────────────────────────────────────────────────

#: Spuren der Standard-Kopf-/Fußzeile eines Browsers. Bewusst eng gefasst:
#: ein Dateipfad als URL und das amerikanische Datum MIT Uhrzeit, wie Chromium
#: es druckt („10/4/26, 2:31 PM“). Ein Bruch wie „3/4“ allein schlägt nicht an.
BROWSER_RAND = (
    (re.compile(r"file://", re.I), "Dateipfad (file://) aus der Fußzeile des Browsers"),
    (re.compile(r"\b\d{1,2}/\d{1,2}/\d{2,4},?\s+\d{1,2}:\d{2}\s*(?:AM|PM)\b", re.I),
     "US-Datum mit Uhrzeit aus der Kopfzeile des Browsers"),
)

#: Platzhalter, die in keinem versandfertigen Dokument stehen dürfen.
PLATZHALTER = (
    "Muster GmbH", "Musterfirma", "Max Mustermann", "Erika Mustermann",
    "Musterstraße 1", "Lorem ipsum", "[Firmenname]", "<Firmenname>",
)

#: A4 in Punkt (1/72 Zoll); Toleranz für Rundung der Erzeuger.
A4_PT = (595.3, 841.9)
_TOLERANZ_PT = 3.0

ANZEIGE_HINWEIS = (
    "Erzeuge das PDF neu mit `dokument pdf <eingabe.md|.html> -o <ziel.pdf>` — "
    "das setzt A4, eine eigene Fußzeile „Seite X von Y“ und keinen Browser-Rand. "
    "Danach erneut present_file aufrufen."
)


def _pdf_text(pfad: str | os.PathLike) -> tuple[list[tuple[float, float]], str]:
    """(Seitengrößen in pt, gesamter Text). pymupdf im Abbild, sonst pypdf."""
    try:
        import fitz  # pymupdf

        with fitz.open(str(pfad)) as doc:
            groessen = [(p.rect.width, p.rect.height) for p in doc]
            text = "\n".join(p.get_text() for p in doc)
        return groessen, text
    except ImportError:
        pass
    from pypdf import PdfReader

    reader = PdfReader(str(pfad))
    groessen = [(float(p.mediabox.width), float(p.mediabox.height)) for p in reader.pages]
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    return groessen, text


def _ist_a4(breite: float, hoehe: float) -> bool:
    b, h = sorted((breite, hoehe))
    return abs(b - A4_PT[0]) <= _TOLERANZ_PT and abs(h - A4_PT[1]) <= _TOLERANZ_PT


def pruefe_pdf(pfad: str | os.PathLike) -> dict:
    """Ist das PDF versandfertig?

    ``fehler`` (je ``art`` + ``text``): ``browser_rand`` oder ``platzhalter``.
    ``hinweise``: z. B. ein anderes Papierformat als A4.
    ``anzeige_meldung``: gesetzt, wenn ``present_file`` das PDF abweisen soll —
    nur bei Browser-Rand; ein Platzhalter kann in einem Entwurf berechtigt sein.
    """
    fehler: list[dict] = []
    hinweise: list[str] = []
    try:
        groessen, text = _pdf_text(pfad)
    except Exception as e:  # noqa: BLE001 — kaputtes PDF ist ein Befund, kein Absturz
        return {
            "ok": False, "seiten": 0, "format": None, "hinweise": [],
            "fehler": [{"art": "unlesbar", "text": f"PDF nicht lesbar: {e}"}],
            "anzeige_meldung": None,
        }

    for muster, beschreibung in BROWSER_RAND:
        treffer = muster.search(text)
        if treffer:
            fehler.append({"art": "browser_rand", "text": f"{beschreibung}: „{treffer.group(0)}“"})

    klein = text.lower()
    for platzhalter in PLATZHALTER:
        if platzhalter.lower() in klein:
            fehler.append({"art": "platzhalter", "text": f"Platzhalter im Text: „{platzhalter}“"})

    alle_a4 = bool(groessen) and all(_ist_a4(b, h) for b, h in groessen)
    if groessen and not alle_a4:
        b, h = groessen[0]
        hinweise.append(f"Papierformat ist nicht A4 ({b:.0f} × {h:.0f} pt).")

    rand = [f["text"] for f in fehler if f["art"] == "browser_rand"]
    meldung = None
    if rand:
        meldung = ("Dieses PDF trägt die Kopf-/Fußzeile eines Browsers ("
                   + "; ".join(rand) + ") und ist so nicht versandfertig. " + ANZEIGE_HINWEIS)
    return {
        "ok": not fehler,
        "seiten": len(groessen),
        "format": "A4" if alle_a4 else ("anderes" if groessen else None),
        "fehler": fehler,
        "hinweise": hinweise,
        "anzeige_meldung": meldung,
    }


# ─── PDF erzeugen ─────────────────────────────────────────────────────────────

DRUCK_CSS = """
@page { size: A4; margin: 20mm 18mm 22mm 18mm; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: "Liberation Sans", "DejaVu Sans", Arial, Helvetica, sans-serif;
       font-size: 10.5pt; line-height: 1.45; color: #1a1a1a; margin: 0; }
h1 { font-size: 18pt; margin: 0 0 10pt; }
h2 { font-size: 14pt; margin: 16pt 0 6pt; }
h3 { font-size: 12pt; margin: 12pt 0 4pt; }
p { margin: 0 0 7pt; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0 10pt; page-break-inside: auto; }
tr { page-break-inside: avoid; }
th, td { border: 0.6pt solid #9a9a9a; padding: 3pt 5pt; vertical-align: top; text-align: left; }
th { background: #eeeeee; }
code, pre { font-family: "DejaVu Sans Mono", "Liberation Mono", monospace; font-size: 9pt; }
pre { white-space: pre-wrap; }
img { max-width: 100%; }
.seitenumbruch { page-break-after: always; }
"""

_FUSSZEILE = (
    '<div style="width:100%;font-size:8pt;color:#555;padding:0 18mm;'
    'font-family:Liberation Sans,DejaVu Sans,Arial,sans-serif;display:flex;'
    'justify-content:space-between;">'
    '<span>{links}</span>'
    '<span>Seite <span class="pageNumber"></span> von <span class="totalPages"></span></span>'
    "</div>"
)


#: Herkunft, unter der das Dokument gerendert wird. Gibt es nicht im Netz
#: (.invalid ist reserviert) — jede Anfrage dorthin beantwortet der Route-
#: Handler selbst aus dem Ordner des Quelldokuments. Kein file://: so greift
#: zusätzlich die Sperre des Browsers gegen lokale Dateien aus einer Webseite.
HERKUNFT = "http://dokument.invalid/"
_SEITE = "__dokument__.html"

#: Elemente, die nachladen, einbetten, weiterleiten oder Code ausführen. Im
#: Dokument eines Agenten haben sie nichts verloren; Eingaben können aus
#: fremden Quellen stammen (Webseiten, Prompt-Injektion).
_GEFAEHRLICH_PAARE = ("script", "iframe", "object", "frame", "frameset", "noscript", "template", "applet")
_GEFAEHRLICH_EINZELN = ("embed", "link", "base", "frame", "iframe", "object", "script", "portal")


def _entschaerfen(text: str) -> str:
    """Nachladende/ausführende Elemente und Weiterleitungen entfernen — bis nichts mehr fällt.

    Wiederholt, weil verschachtelte Eingaben wie ``<ifr<iframe>ame …>`` nach
    einem Durchgang wieder ein gültiges Element ergeben. Zweite Linie: die
    erste ist der Route-Handler beim Rendern (nur Bilder/Schriften/Stylesheets
    aus dem Ordner), plus abgeschaltetes JavaScript.
    """
    for _ in range(20):
        neu = _entschaerfen_einmal(text)
        if neu == text:
            return neu
        text = neu
    # Hört nicht auf zu schrumpfen: lieber gar kein Markup als ein halbes.
    return html.escape(text)


def _entschaerfen_einmal(text: str) -> str:
    for tag in _GEFAEHRLICH_PAARE:
        text = re.sub(rf"<\s*{tag}\b.*?<\s*/\s*{tag}\s*>", "", text, flags=re.I | re.S)
    for tag in _GEFAEHRLICH_EINZELN:
        text = re.sub(rf"<\s*/?\s*{tag}\b[^>]*>", "", text, flags=re.I)
    # Weiterleitung per <meta http-equiv="refresh">
    text = re.sub(r"<\s*meta\b[^>]*http-equiv[^>]*>", "", text, flags=re.I)
    # Ereignis-Attribute (onload=…) und javascript:-Adressen
    text = re.sub(r"\son[a-z]+\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", "", text, flags=re.I)
    text = re.sub(r"javascript\s*:", "", text, flags=re.I)
    return text


def _im_ordner(basis: Path, relativ: str) -> Path | None:
    """Datei unterhalb von ``basis`` — nach Auflösung aller Symlinks — oder None."""
    if not relativ or "\x00" in relativ:
        return None
    wurzel = Path(os.path.realpath(basis))
    kandidat = Path(os.path.realpath(wurzel / relativ))
    try:
        kandidat.relative_to(wurzel)
    except ValueError:
        return None
    return kandidat if kandidat.is_file() else None


#: Was beim Rendern aus dem Ordner nachgeladen werden darf: nur Unterressourcen,
#: die sichtbar ins PDF gehen, aber keinen Dateiinhalt als Text zeigen können.
#: Frames/Dokumente (iframe, object) und Textdateien bleiben draußen — sonst
#: würde ein eingeschleustes ``<iframe src="notizen.txt">`` sie ins PDF drucken.
_ERLAUBTE_TYPEN = frozenset({"image", "font", "stylesheet"})
_ERLAUBTE_ENDUNGEN = frozenset({
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".bmp", ".ico",
    ".woff", ".woff2", ".ttf", ".otf", ".css",
})


def _ressource_aus_ordner(basis: Path, url: str, typ: str) -> Path | None:
    """Datei für eine Nachlade-Anfrage beim Rendern — oder None (= blockieren)."""
    from urllib.parse import unquote, urlsplit

    if typ not in _ERLAUBTE_TYPEN or not url.startswith(HERKUNFT):
        return None
    teile = urlsplit(url)
    if teile.query or teile.fragment:
        return None
    datei = _im_ordner(basis, unquote(teile.path).lstrip("/"))
    if datei is None or datei.suffix.lower() not in _ERLAUBTE_ENDUNGEN:
        return None
    # Stylesheet nur als .css, Bild/Schrift nie als .css
    if (typ == "stylesheet") != (datei.suffix.lower() == ".css"):
        return None
    return datei


def _markdown_zu_html(text: str) -> str:
    import markdown

    return markdown.markdown(
        text, extensions=["tables", "fenced_code", "sane_lists", "attr_list"],
        output_format="html5",
    )


def _erste_ueberschrift(html_text: str) -> str | None:
    treffer = re.search(r"<h1[^>]*>(.*?)</h1>", html_text, re.S | re.I)
    if not treffer:
        return None
    return html.unescape(re.sub(r"<[^>]+>", "", treffer.group(1))).strip() or None


def baue_html(eingabe: Path, titel: str | None = None) -> str:
    """Vollständiges HTML-Dokument (lang=de, Druck-CSS) aus Markdown oder HTML.

    Relative Bilder/Stylesheets lösen gegen ``HERKUNFT`` auf; der Route-Handler
    in ``erzeuge_pdf`` liefert sie aus dem Ordner des Quelldokuments.
    """
    roh = eingabe.read_text(encoding="utf-8")
    basis = f'<base href="{HERKUNFT}">'
    if eingabe.suffix.lower() in (".md", ".markdown", ".txt"):
        # Markdown reicht Roh-HTML durch — deshalb erst NACH der Umwandlung entschärfen.
        koerper = _entschaerfen(_markdown_zu_html(roh))
        titel = titel or _erste_ueberschrift(koerper) or eingabe.stem
        return (
            '<!doctype html><html lang="de"><head><meta charset="utf-8">'
            f"{basis}<title>{html.escape(titel)}</title>"
            f"<style>{DRUCK_CSS}</style></head><body>{koerper}</body></html>"
        )
    # HTML: eigenes CSS des Agenten bleibt maßgeblich — unseres kommt ZUERST,
    # damit es nur füllt, was dort fehlt (A4, Schrift, Tabellenränder).
    roh = _entschaerfen(roh)
    kopf = f'<meta charset="utf-8">{basis}<style>{DRUCK_CSS}</style>'
    if titel:
        kopf += f"<title>{html.escape(titel)}</title>"
    if re.search(r"<head[^>]*>", roh, re.I):
        dokument_html = re.sub(r"(<head[^>]*>)", lambda m: m.group(1) + kopf, roh, count=1, flags=re.I)
    elif re.search(r"<html[^>]*>", roh, re.I):
        dokument_html = re.sub(r"(<html[^>]*>)", lambda m: m.group(1) + "<head>" + kopf + "</head>",
                               roh, count=1, flags=re.I)
    else:
        dokument_html = f'<!doctype html><html lang="de"><head>{kopf}</head><body>{roh}</body></html>'
    if not re.search(r"<html[^>]*\blang=", dokument_html, re.I):
        dokument_html = re.sub(r"<html", '<html lang="de"', dokument_html, count=1, flags=re.I)
    return dokument_html


def _chromium_pfad() -> str | None:
    """Dieselbe Browserwahl wie das Browser-Werkzeug des Agenten."""
    try:
        from app.tools.browser import _chromium_pfad as browser_pfad

        return browser_pfad()
    except Exception:  # noqa: BLE001 — außerhalb des Abbilds: Playwright-Standard
        return None


def erzeuge_pdf(eingabe: Path, ziel: Path, titel: str | None = None, fusszeile: str = "") -> dict:
    """PDF rendern — abgeschottet. Rückgabe: welche Anfragen erlaubt/blockiert wurden.

    Die Eingabe kann aus fremden Quellen stammen (Webinhalte, Prompt-Injektion),
    und der Agenten-Container erreicht das interne Netz. Deshalb:

    * kein JavaScript, keine Service Worker;
    * JEDE Anfrage läuft durch den Route-Handler: beantwortet wird nur die Seite
      selbst und Dateien INNERHALB des Ordners des Quelldokuments (nach realpath,
      also auch kein Symlink nach draußen, kein ``..``). Alles andere — http(s),
      file://, andere Schemata — wird abgebrochen, bevor es das Netz erreicht;
    * nachladende Elemente (iframe, object, embed, link, script, Weiterleitung)
      werden vorher entfernt.
    """
    from playwright.sync_api import sync_playwright

    seite_html = baue_html(eingabe, titel)
    basis = eingabe.resolve().parent
    bericht: dict[str, list[str]] = {"erlaubt": [], "blockiert": []}

    def weiche(route):
        url = route.request.url
        if url == HERKUNFT + _SEITE:
            route.fulfill(status=200, content_type="text/html; charset=utf-8", body=seite_html)
            return
        datei = _ressource_aus_ordner(basis, url, route.request.resource_type)
        if datei is not None:
            bericht["erlaubt"].append(url)
            route.fulfill(path=str(datei))
            return
        bericht["blockiert"].append(url)
        route.abort("blockedbyclient")

    ziel.parent.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as pw:
        pfad = _chromium_pfad()
        browser = pw.chromium.launch(
            args=["--no-sandbox", "--disable-dev-shm-usage"],
            **({"executable_path": pfad} if pfad else {}),
        )
        try:
            kontext = browser.new_context(
                locale="de-DE", java_script_enabled=False, service_workers="block",
            )
            kontext.route("**/*", weiche)
            seite = kontext.new_page()
            seite.goto(HERKUNFT + _SEITE, wait_until="load")
            seite.pdf(
                path=str(ziel),
                format="A4",
                print_background=True,
                display_header_footer=True,
                # Leere Kopfzeile statt der Standardzeile (Datum + Titel).
                header_template="<span></span>",
                footer_template=_FUSSZEILE.format(links=html.escape(fusszeile)),
                margin={"top": "20mm", "bottom": "22mm", "left": "18mm", "right": "18mm"},
            )
        finally:
            browser.close()
    return bericht


# ─── Word erzeugen ────────────────────────────────────────────────────────────

class _DocxBauer(HTMLParser):
    """Wandelt das HTML aus ``markdown`` in ein Word-Dokument um.

    Abgedeckt ist, was Markdown erzeugt: Überschriften, Absätze, Listen,
    Tabellen, fett/kursiv/Code, Zeilenumbruch, Trennlinie.
    """

    def __init__(self, doc):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.absatz = None
        self.fett = self.kursiv = self.code = 0
        self.listen: list[str] = []
        self.tabelle_zeilen: list[list[str]] | None = None
        self.zelle: list[str] | None = None
        self.kopfzeile = False

    def _neuer_absatz(self, stil: str | None = None):
        self.absatz = self.doc.add_paragraph(style=stil) if stil else self.doc.add_paragraph()

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.absatz = self.doc.add_heading(level=min(int(tag[1]), 4))
        elif tag == "p" and self.zelle is None:
            if not self.listen:
                self._neuer_absatz()
        elif tag in ("ul", "ol"):
            self.listen.append(tag)
        elif tag == "li":
            self._neuer_absatz("List Number" if self.listen and self.listen[-1] == "ol" else "List Bullet")
        elif tag in ("strong", "b"):
            self.fett += 1
        elif tag in ("em", "i"):
            self.kursiv += 1
        elif tag == "code":
            self.code += 1
        elif tag == "br":
            if self.zelle is not None:
                self.zelle.append("\n")
            elif self.absatz is not None:
                self.absatz.add_run().add_break()
        elif tag == "table":
            self.tabelle_zeilen = []
        elif tag == "tr" and self.tabelle_zeilen is not None:
            self.tabelle_zeilen.append([])
        elif tag in ("td", "th"):
            self.zelle = []
        elif tag == "hr":
            self._neuer_absatz()
        elif tag == "pre":
            self._neuer_absatz()

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre"):
            if tag != "p" or self.zelle is None:
                self.absatz = None
        elif tag in ("ul", "ol") and self.listen:
            self.listen.pop()
        elif tag in ("strong", "b"):
            self.fett = max(0, self.fett - 1)
        elif tag in ("em", "i"):
            self.kursiv = max(0, self.kursiv - 1)
        elif tag == "code":
            self.code = max(0, self.code - 1)
        elif tag in ("td", "th") and self.zelle is not None and self.tabelle_zeilen:
            self.tabelle_zeilen[-1].append("".join(self.zelle).strip())
            self.zelle = None
        elif tag == "table" and self.tabelle_zeilen is not None:
            self._tabelle_schreiben(self.tabelle_zeilen)
            self.tabelle_zeilen = None

    def handle_data(self, data):
        if self.zelle is not None:
            self.zelle.append(data)
            return
        if not data.strip() and self.absatz is None:
            return
        if self.absatz is None:
            self._neuer_absatz()
        run = self.absatz.add_run(data)
        run.bold = bool(self.fett) or None
        run.italic = bool(self.kursiv) or None
        if self.code:
            run.font.name = "DejaVu Sans Mono"

    def _tabelle_schreiben(self, zeilen):
        zeilen = [z for z in zeilen if z]
        if not zeilen:
            return
        spalten = max(len(z) for z in zeilen)
        tabelle = self.doc.add_table(rows=len(zeilen), cols=spalten)
        try:
            tabelle.style = "Table Grid"
        except Exception:  # noqa: BLE001 — Vorlage ohne diesen Stil
            pass
        for r, zeile in enumerate(zeilen):
            for c, wert in enumerate(zeile):
                zelle = tabelle.cell(r, c)
                zelle.text = wert
                if r == 0:
                    for run in zelle.paragraphs[0].runs:
                        run.bold = True


def erzeuge_docx(eingabe: Path, ziel: Path) -> None:
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Mm

    roh = eingabe.read_text(encoding="utf-8")
    inhalt = roh if eingabe.suffix.lower() in (".html", ".htm") else _markdown_zu_html(roh)
    doc = docx.Document()
    for abschnitt in doc.sections:
        abschnitt.page_width, abschnitt.page_height = Mm(210), Mm(297)
        abschnitt.left_margin = abschnitt.right_margin = Mm(20)
        abschnitt.top_margin = abschnitt.bottom_margin = Mm(20)
    # Sprache Deutsch für Silbentrennung und Rechtschreibprüfung in Word.
    stil = doc.styles["Normal"]
    rpr = stil.element.get_or_add_rPr()
    sprache = rpr.find(qn("w:lang"))
    if sprache is None:
        sprache = rpr.makeelement(qn("w:lang"), {})
        rpr.append(sprache)
    sprache.set(qn("w:val"), "de-DE")
    doc.core_properties.language = "de-DE"
    doc.core_properties.title = _erste_ueberschrift(inhalt) or eingabe.stem

    _DocxBauer(doc).feed(inhalt)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ziel))


# ─── Kommandozeile ────────────────────────────────────────────────────────────

def _bericht(ergebnis: dict) -> str:
    zeilen = [f"{ergebnis['seiten']} Seite(n), Format: {ergebnis['format'] or 'unbekannt'}"]
    zeilen += [f"FEHLER ({f['art']}): {f['text']}" for f in ergebnis["fehler"]]
    zeilen += [f"Hinweis: {h}" for h in ergebnis["hinweise"]]
    zeilen.append("versandfertig" if ergebnis["ok"] else "NICHT versandfertig")
    return "\n".join(zeilen)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="dokument",
        description="Versandfertige PDF- und Word-Dokumente aus Markdown/HTML erzeugen und prüfen.",
    )
    befehle = parser.add_subparsers(dest="befehl", required=True)

    p_pdf = befehle.add_parser("pdf", help="PDF (A4, Fußzeile „Seite X von Y“, ohne Browser-Rand)")
    p_pdf.add_argument("eingabe")
    p_pdf.add_argument("-o", "--ausgabe")
    p_pdf.add_argument("--titel")
    p_pdf.add_argument("--fusszeile", default="", help="Text links in der Fußzeile, z. B. Firmenname")

    p_docx = befehle.add_parser("docx", help="Word-Dokument (A4, Deutsch) aus Markdown")
    p_docx.add_argument("eingabe")
    p_docx.add_argument("-o", "--ausgabe")

    p_pruefen = befehle.add_parser("pruefen", help="PDF auf Browser-Rand, Platzhalter und A4 prüfen")
    p_pruefen.add_argument("datei")
    p_pruefen.add_argument("--json", action="store_true")

    args = parser.parse_args(argv)

    if args.befehl == "pruefen":
        ergebnis = pruefe_pdf(args.datei)
        print(json.dumps(ergebnis, ensure_ascii=False) if args.json else _bericht(ergebnis))
        return 0 if ergebnis["ok"] else 1

    eingabe = Path(args.eingabe)
    if not eingabe.is_file():
        print(f"Eingabe nicht gefunden: {eingabe}", file=sys.stderr)
        return 2

    if args.befehl == "docx":
        ziel = Path(args.ausgabe) if args.ausgabe else eingabe.with_suffix(".docx")
        erzeuge_docx(eingabe, ziel)
        print(f"Word-Dokument erstellt: {ziel}")
        return 0

    ziel = Path(args.ausgabe) if args.ausgabe else eingabe.with_suffix(".pdf")
    if eingabe.suffix.lower() not in (".md", ".markdown", ".txt", ".html", ".htm"):
        print("Eingabe muss Markdown (.md) oder HTML (.html) sein.", file=sys.stderr)
        return 2
    erzeuge_pdf(eingabe, ziel, titel=args.titel, fusszeile=args.fusszeile)
    ergebnis = pruefe_pdf(ziel)
    print(f"PDF erstellt: {ziel}")
    print(_bericht(ergebnis))
    return 0 if ergebnis["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
