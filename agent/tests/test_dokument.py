"""Versandfertige Dokumente: ``dokument pdf`` statt Browser-Druck (#893).

Befund aus dem Markttest: Vom Agenten erzeugte PDFs trugen die Kopf- und
Fußzeile des Browsers — oben ein US-Datum („10/4/26, 2:31 PM“), unten den
Dateipfad (``file:///workspace/…``). Das Abbild hatte kein PDF-Werkzeug, also
improvisierten die Agenten mit ``chromium --print-to-pdf``.

Jetzt gibt es EIN Werkzeug für alle Laufzeiten (``dokument``), und
``present_file`` weist ein PDF mit Browser-Rand ab — mit dem Hinweis, wie es
richtig geht. Die Prüfung steckt in ``app.dokument``; der MCP-Server (Claude
Code) ruft sie über die Kommandozeile, der Executor (Codex/Custom-LLM) direkt.
"""

import asyncio
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from app import dokument

AGENT = Path(__file__).resolve().parents[1]
CLI = AGENT / "scripts" / "dokument"


def _mini_pdf(zeilen: list[str], breite: int = 595, hoehe: int = 842) -> bytes:
    """Ein einseitiges PDF mit Textzeilen — ohne Browser, ohne Bibliothek."""
    def esc(t: str) -> str:
        return t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    inhalt = "BT /F1 10 Tf 40 {} Td 14 TL ".format(hoehe - 40)
    inhalt += " ".join(f"({esc(z)}) Tj T*" for z in zeilen) + " ET"
    objekte = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {breite} {hoehe}] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(inhalt.encode('latin-1'))} >>\nstream\n{inhalt}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    aus = b"%PDF-1.4\n"
    versatz = []
    for nr, obj in enumerate(objekte, start=1):
        versatz.append(len(aus))
        aus += f"{nr} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(aus)
    aus += f"xref\n0 {len(objekte) + 1}\n0000000000 65535 f \n".encode()
    for v in versatz:
        aus += f"{v:010d} 00000 n \n".encode()
    aus += f"trailer\n<< /Size {len(objekte) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return aus


class _MitOrdner(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.ordner = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _pdf(self, name: str, zeilen: list[str], **kw) -> Path:
        pfad = self.ordner / name
        pfad.write_bytes(_mini_pdf(zeilen, **kw))
        return pfad


BROWSER_DRUCK = [
    "10/4/26, 2:31 PM                         Angebot",
    "Angebot Nr. 17",
    "file:///workspace/transfer/angebot.html                  1/1",
]


class PruefenTests(_MitOrdner):
    def test_browser_kopf_und_fusszeile_werden_erkannt(self):
        ergebnis = dokument.pruefe_pdf(self._pdf("roh.pdf", BROWSER_DRUCK))
        arten = {f["art"] for f in ergebnis["fehler"]}
        self.assertIn("browser_rand", arten)
        self.assertFalse(ergebnis["ok"])
        self.assertIn("dokument pdf", ergebnis["anzeige_meldung"])

    def test_sauberes_a4_dokument_besteht(self):
        ergebnis = dokument.pruefe_pdf(self._pdf("gut.pdf", ["Angebot Nr. 17", "Seite 1 von 1"]))
        self.assertTrue(ergebnis["ok"], ergebnis)
        self.assertEqual(ergebnis["fehler"], [])
        self.assertEqual(ergebnis["format"], "A4")
        self.assertIsNone(ergebnis["anzeige_meldung"])

    def test_platzhalter_sind_ein_fehler_aber_kein_grund_zur_ablehnung_beim_anzeigen(self):
        ergebnis = dokument.pruefe_pdf(self._pdf("muster.pdf", ["Kunde: Muster GmbH"]))
        self.assertFalse(ergebnis["ok"])
        self.assertEqual({f["art"] for f in ergebnis["fehler"]}, {"platzhalter"})
        self.assertIsNone(ergebnis["anzeige_meldung"])

    def test_anderes_papierformat_ist_ein_hinweis(self):
        ergebnis = dokument.pruefe_pdf(self._pdf("letter.pdf", ["Text"], breite=612, hoehe=792))
        self.assertTrue(ergebnis["ok"])
        self.assertTrue(any("A4" in h for h in ergebnis["hinweise"]))

    def test_kommandozeile_meldet_fehler_ueber_den_exitcode(self):
        roh = self._pdf("roh.pdf", BROWSER_DRUCK)
        lauf = subprocess.run([sys.executable, str(CLI), "pruefen", "--json", str(roh)],
                              capture_output=True, text=True, cwd=AGENT)
        self.assertEqual(lauf.returncode, 1, lauf.stderr)
        self.assertIn("browser_rand", {f["art"] for f in json.loads(lauf.stdout)["fehler"]})


class PresentFileTests(_MitOrdner):
    def _executor(self):
        from app.tools.executor import ToolExecutor

        ex = ToolExecutor.__new__(ToolExecutor)
        ex._resolve_path = lambda p: p  # noqa: SLF001 — Pfadauflösung ist hier nicht Thema
        return ex

    def _zeige(self, pfad: Path) -> str:
        return asyncio.run(self._executor()._tool_present_file({"path": str(pfad)}))

    def test_pdf_mit_browser_rand_wird_abgewiesen(self):
        antwort = self._zeige(self._pdf("roh.pdf", BROWSER_DRUCK))
        self.assertTrue(antwort.startswith("Error"), antwort)
        self.assertIn("dokument pdf", antwort)

    def test_sauberes_pdf_wird_angezeigt(self):
        antwort = self._zeige(self._pdf("gut.pdf", ["Angebot", "Seite 1 von 1"]))
        self.assertTrue(antwort.startswith("__AI_EMPLOYEE_PRESENT_FILE__"), antwort)

    def test_andere_dateien_bleiben_unberuehrt(self):
        datei = self.ordner / "notiz.txt"
        datei.write_text("file:///workspace/x 10/4/26, 2:31 PM")
        antwort = self._zeige(datei)
        self.assertTrue(antwort.startswith("__AI_EMPLOYEE_PRESENT_FILE__"), antwort)


MARKDOWN = """# Angebot 2026-017

Sehr geehrte Damen und Herren,

vielen Dank für Ihre Anfrage. Wir bieten an:

| Pos. | Leistung | Menge | Preis |
|---|---|---|---|
| 1 | Wand streichen | 42,50 m² | 510,00 € |

- Gültig 30 Tage
- Zahlbar netto 14 Tage
"""


def _browser_verfuegbar() -> bool:
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as pw:
            pfad = dokument._chromium_pfad()
            browser = pw.chromium.launch(**({"executable_path": pfad} if pfad else {}))
            browser.close()
        return True
    except Exception:  # noqa: BLE001
        return False


class PdfErzeugenTests(_MitOrdner):
    @classmethod
    def setUpClass(cls):
        if not _browser_verfuegbar():
            raise unittest.SkipTest("Kein Chromium für Playwright verfügbar")

    def test_markdown_wird_zu_a4_ohne_browser_rand_mit_seitenzahl(self):
        quelle = self.ordner / "angebot.md"
        quelle.write_text(MARKDOWN, encoding="utf-8")
        ziel = self.ordner / "angebot.pdf"
        lauf = subprocess.run([sys.executable, str(CLI), "pdf", str(quelle), "-o", str(ziel)],
                              capture_output=True, text=True, cwd=AGENT)
        self.assertEqual(lauf.returncode, 0, lauf.stdout + lauf.stderr)
        ergebnis = dokument.pruefe_pdf(ziel)
        self.assertTrue(ergebnis["ok"], ergebnis)
        self.assertEqual(ergebnis["format"], "A4")
        text = dokument._pdf_text(ziel)[1]
        self.assertIn("Seite 1 von 1", text)
        self.assertIn("Wand streichen", text)
        self.assertNotIn("file://", text)


class DocxTests(_MitOrdner):
    def test_markdown_wird_zu_word_mit_ueberschrift_tabelle_und_liste(self):
        import docx

        quelle = self.ordner / "angebot.md"
        quelle.write_text(MARKDOWN, encoding="utf-8")
        ziel = self.ordner / "angebot.docx"
        lauf = subprocess.run([sys.executable, str(CLI), "docx", str(quelle), "-o", str(ziel)],
                              capture_output=True, text=True, cwd=AGENT)
        self.assertEqual(lauf.returncode, 0, lauf.stdout + lauf.stderr)
        doc = docx.Document(str(ziel))
        absaetze = [p for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(absaetze[0].text, "Angebot 2026-017")
        self.assertTrue(absaetze[0].style.name.startswith("Heading"))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(1, 1).text, "Wand streichen")
        listen = [p for p in doc.paragraphs if p.style.name.startswith("List")]
        self.assertEqual([p.text for p in listen], ["Gültig 30 Tage", "Zahlbar netto 14 Tage"])
        abschnitt = doc.sections[0]
        self.assertAlmostEqual(abschnitt.page_width.mm, 210, delta=1)
        self.assertAlmostEqual(abschnitt.page_height.mm, 297, delta=1)


class AnleitungTests(unittest.TestCase):
    """Die Regel „PDF nur mit `dokument pdf`“ erreicht jede Laufzeit und jeden Kanal."""

    REGEL = "dokument pdf"

    def test_fertige_ergebnisse_in_der_gemeinsamen_anleitung(self):
        from app.runner_hooks import MULTIMODAL_CAPABILITY_NOTE

        self.assertIn(self.REGEL, MULTIMODAL_CAPABILITY_NOTE)
        self.assertIn("--print-to-pdf", MULTIMODAL_CAPABILITY_NOTE)

    def test_chat_und_telegram_bekommen_die_regel(self):
        from app import chat_consumer

        kanal = chat_consumer._build_channel_prompt("Schick mir das Angebot als PDF", "web", True)
        telegram = chat_consumer._build_telegram_prompt("Schick das PDF", {"chat_id": 1}, True)
        self.assertIn(self.REGEL, kanal)
        self.assertIn(self.REGEL, telegram)

    def test_claude_code_anleitung(self):
        self.assertIn(self.REGEL, (AGENT / "claude-global.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    os.environ.setdefault("PYTHONPATH", str(AGENT))
    unittest.main()
