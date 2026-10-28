"""Dateien am Projekt (Wissen → Dateien).

Original zum Herunterladen neben der Text-Ablage, ausgelesener Text als Seite im
Projektwissen (durchsuchbar fuer Agenten). Geprueft ueber die echten Endpunkte:
Auslesen (Word, Excel, Text), Namen und Pfade, Groesse und Art, Herunterladen nur
als Anhang, Loeschen samt Seite, Betrachter und fremde Projekte.
"""

import io
import os

from fastapi import HTTPException
from starlette.datastructures import UploadFile

from app.api import teams as api
from app.core import projekt_dateien, vault
from tests.test_projekte import BETRACHTER, MITGLIED, Basis


def _datei(name: str, inhalt: bytes) -> UploadFile:
    return UploadFile(file=io.BytesIO(inhalt), filename=name)


def _docx(text: str) -> bytes:
    from docx import Document

    doc, puffer = Document(), io.BytesIO()
    doc.add_paragraph(text)
    doc.save(puffer)
    return puffer.getvalue()


def _xlsx(zeilen: list[list]) -> bytes:
    from openpyxl import Workbook

    wb, puffer = Workbook(), io.BytesIO()
    for z in zeilen:
        wb.active.append(z)
    wb.save(puffer)
    return puffer.getvalue()


class Dateien(Basis):
    async def _hochladen(self, name, inhalt, user=MITGLIED, projekt="p1"):
        async with self.Session() as db:
            return await api.datei_hochladen(projekt, _datei(name, inhalt), user=user, db=db)

    async def _brain(self):
        from app.core import projekt_wissen
        from app.models.team import Team

        async with self.Session() as db:
            return await projekt_wissen.brain_des_projekts(db, await db.get(Team, "p1"))

    async def test_word_und_excel_werden_ausgelesen_und_als_seite_abgelegt(self):
        r = await self._hochladen("Lastenheft.docx", _docx("Die Anlage muss 400 Teile je Stunde schaffen."))
        self.assertEqual((r["name"], r["seite"], r["text_ausgelesen"]),
                         ("Lastenheft.docx", "dateien/Lastenheft.docx.md", True))
        await self._hochladen("Angebot.xlsx", _xlsx([["Position", "Preis"], ["Analyse", 4200]]))
        brain = await self._brain()
        self.assertIn("400 Teile je Stunde", vault.read_file(brain.host_path, "dateien/Lastenheft.docx.md"))
        self.assertIn("Analyse\t4200", vault.read_file(brain.host_path, "dateien/Angebot.xlsx.md"))
        async with self.Session() as db:
            liste = await api.dateien_liste("p1", user=MITGLIED, db=db)
        self.assertEqual([d["name"] for d in liste["dateien"]], ["Angebot.xlsx", "Lastenheft.docx"])
        # Das Original liegt NICHT in der git-Ablage, sondern daneben.
        self.assertFalse(os.path.exists(os.path.join(brain.host_path, "Lastenheft.docx")))
        self.assertTrue(os.path.isfile(brain.host_path + ".dateien/Lastenheft.docx"))

    async def test_namen_werden_bereinigt_und_nicht_ueberschrieben(self):
        r1 = await self._hochladen("../../etc/notiz.txt", b"eins")
        r2 = await self._hochladen("notiz.txt", b"zwei")
        self.assertEqual((r1["name"], r2["name"]), ("notiz.txt", "notiz (2).txt"))
        r3 = await self._hochladen("a/b\\..\\geheim<script>.md", b"x")
        self.assertEqual(r3["name"], "geheim_script_.md")

    async def test_art_groesse_und_leere_datei(self):
        for name, inhalt, status in (("boese.exe", b"MZ", 400), ("seite.html", b"<html>", 400),
                                     (".versteckt", b"x", 400), ("leer.txt", b"", 400)):
            with self.subTest(name=name), self.assertRaises(HTTPException) as e:
                await self._hochladen(name, inhalt)
            self.assertEqual(e.exception.status_code, status)
        with self.assertRaises(HTTPException) as e:
            await self._hochladen("gross.txt", b"x" * (projekt_dateien.MAX_BYTES + 1))
        self.assertEqual(e.exception.status_code, 413)

    async def test_herunterladen_nur_als_anhang_und_nur_vorhandenes(self):
        await self._hochladen("bericht.txt", b"Inhalt")
        async with self.Session() as db:
            antwort = await api.datei_herunterladen("p1", "bericht.txt", user=MITGLIED, db=db)
            self.assertEqual(antwort.media_type, "application/octet-stream")
            self.assertIn("attachment", antwort.headers["content-disposition"])
            for falsch in ("../bericht.txt.md", "gibtsnicht.txt", "..%2Fx"):
                with self.subTest(falsch=falsch), self.assertRaises(HTTPException) as e:
                    await api.datei_herunterladen("p1", falsch, user=MITGLIED, db=db)
                self.assertEqual(e.exception.status_code, 404)

    async def test_loeschen_entfernt_original_und_seite(self):
        await self._hochladen("weg.txt", b"bald weg")
        async with self.Session() as db:
            await api.datei_loeschen("p1", "weg.txt", user=MITGLIED, db=db)
        brain = await self._brain()
        self.assertFalse(os.path.exists(brain.host_path + ".dateien/weg.txt"))
        with self.assertRaises(FileNotFoundError):
            vault.read_file(brain.host_path, "dateien/weg.txt.md")

    async def test_betrachter_und_fremdes_projekt(self):
        with self.assertRaises(HTTPException) as e:
            await self._hochladen("x.txt", b"x", user=BETRACHTER)
        self.assertEqual(e.exception.status_code, 403)
        with self.assertRaises(HTTPException) as e:
            await self._hochladen("x.txt", b"x", projekt="p2")
        self.assertEqual(e.exception.status_code, 403)
        async with self.Session() as db:
            with self.assertRaises(HTTPException) as e:
                await api.dateien_liste("p2", user=MITGLIED, db=db)
        self.assertEqual(e.exception.status_code, 403)
